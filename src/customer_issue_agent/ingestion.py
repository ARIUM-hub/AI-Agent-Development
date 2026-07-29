from __future__ import annotations

import csv
from io import BytesIO, StringIO
from pathlib import Path

from openpyxl import load_workbook


def extract_conversation_text(filename: str, content: bytes) -> str:
    suffix = Path(filename).suffix.lower()

    if suffix in {".txt", ".log"}:
        text = _decode_text(content)
    elif suffix == ".csv":
        text = _extract_csv(content)
    elif suffix in {".xlsx", ".xlsm"}:
        text = _extract_xlsx(content)
    else:
        raise ValueError("仅支持 txt、log、csv、xlsx 或 xlsm 文件")

    cleaned = text.strip()
    if not cleaned:
        raise ValueError("没有可分析内容")
    return cleaned


def _decode_text(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("文件编码无法识别，请使用 UTF-8 导出")


def _extract_csv(content: bytes) -> str:
    decoded = _decode_text(content)
    reader = csv.DictReader(StringIO(decoded))
    rows: list[str] = []
    for row in reader:
        speaker = _first_value(row, ("speaker", "role", "sender", "from", "角色", "发送方")) or "Unknown"
        message = _first_value(row, ("message", "text", "content", "body", "消息", "内容"))
        if message:
            rows.append(f"{speaker}: {message}")
    if rows:
        return "\n".join(rows)
    return decoded


def _extract_xlsx(content: bytes) -> str:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return ""

    headers = [str(value).strip() if value is not None else "" for value in rows[0]]
    speaker_index = _first_index(headers, ("speaker", "role", "sender", "from", "角色", "发送方"))
    message_index = _first_index(headers, ("message", "text", "content", "body", "消息", "内容"))
    lines: list[str] = []

    for row in rows[1:]:
        values = ["" if value is None else str(value).strip() for value in row]
        if message_index is not None and message_index < len(values):
            speaker = values[speaker_index] if speaker_index is not None and speaker_index < len(values) else "Unknown"
            message = values[message_index]
            if message:
                lines.append(f"{speaker}: {message}")
        else:
            line = " ".join(value for value in values if value)
            if line:
                lines.append(line)
    return "\n".join(lines)


def _first_value(row: dict[str, str | None], keys: tuple[str, ...]) -> str | None:
    normalized = {key.strip().lower(): value for key, value in row.items() if key}
    for key in keys:
        value = normalized.get(key.lower())
        if value and value.strip():
            return value.strip()
    return None


def _first_index(headers: list[str], keys: tuple[str, ...]) -> int | None:
    lowered = [header.lower() for header in headers]
    for key in keys:
        if key.lower() in lowered:
            return lowered.index(key.lower())
    return None
